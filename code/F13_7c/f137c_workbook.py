"""Phase F13.7c - step 1: supplementary/Supplementary_tables.xlsx and tables/X2_sheet_sources.csv, built by F13.7a's workbook program
(f137a_workbook.derive / build / deterministic, unchanged) with the extended source list of f137c_lib.sheets_v2: S5c = the v3.1 changes
(F13.2e T1), the former S5c becomes S5d; S6d-S6e (T2f2_v31, T2g_v31); S13 = the citation structure by work key (F13.3b); S14 = the data of
the article's figures in the order of the F13.6c draft; the Phase log to F13.8. Three S13 sheets are derived here under R13's rule for
listed works (first author's surname, year, title of at most 15 words; no corpus author named): T2 and T5b with the first author withheld
where the work is a corpus record or the surname is a corpus author's (checked in memory against E1's Authors), T3c without its record
identifier. The README sheet is rewritten for v2 (contents, sources, date, the licence sentence)."""
import sys
sys.dont_write_bytecode = True

import io
import json
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import f137c_lib as L

WB = L.WB
CONTENTS = [
    ("S1", "queries, retrieval dates and totals"), ("S2", "Scopus Analyze year series"), ("S3", "activity index by period, window and country"),
    ("S4", "keyword thesaurus v4 and keyword layers"), ("S5", "institutions thesaurus v3.1, its changes and the v3 rules"),
    ("S6", "Kazakhstan institutions, concentration, pairs and collaboration types (v3.1)"), ("S7", "countries, partners and international co-authorship"),
    ("S8", "document types, sources, citations, open access and funding"), ("S9", "co-word map of author keywords and its stability"),
    ("S10", "coverage at an equal query, Scopus and OpenAlex"), ("S11", "references and their age (F13.3)"),
    ("S12", "Web of Science validation"), ("S13", "citation structure by work key (F13.3b)"), ("S14", "data of the article's figures (figure_k sheets)"),
    ("Phase_log", "phases F12.4 to F13.8 from the project registry")]


def group(sheet):
    if sheet == "Phase_log":
        return "Phase_log"
    if sheet.startswith("figure_"):
        return "S14"
    return sheet.split("_")[0].rstrip("abcdefghijklmnopqrstuvwxyz")


def derive_v2():
    WB.derive()  # S1_queries, S6c_pairs_v31_top30, S8s_top20_cited (F13.7a's rules)
    t2 = L.read_csv(L.t("phase_F13_3b", "T2_top30_works"))
    corpus_keys = {(r["first_author"], r["year"]) for r in t2 if r["corpus_record"] == "yes"}
    stats = {}

    def cited(rows, flag):
        out, n = [], 0
        for r in rows:
            w = (r.get("corpus_record") == "yes") if flag else ((r["first_author"], r["year"]) in corpus_keys)
            if w or L.is_corpus_author(r["first_author"]):
                n += 1
                r = {**r, "first_author": L.WITHHELD}
            out.append({**r, "title": L.words(r["title"], 15)})
        return out, n

    rows, stats["S13d_withheld"] = cited(t2, True)
    L.write_csv(L.TABLES / "S13d_top30_works.csv", list(t2[0].keys()), rows)
    t5 = L.read_csv(L.t("phase_F13_3b", "T5b_cocitation_clusters"))
    rows, stats["S13m_withheld"] = cited(t5, False)
    L.write_csv(L.TABLES / "S13m_cocitation_clusters.csv", list(t5[0].keys()), rows)
    t3 = L.read_csv(L.t("phase_F13_3b", "T3c_top20_LCS"))
    keep = [k for k in t3[0].keys() if k != "doi_or_eid"]
    L.write_csv(L.TABLES / "S13h_top20_LCS.csv", keep, [{**{k: r[k] for k in keep}, "title": L.words(r["title"], 15)} for r in t3])
    stats.update({"S13d_rows": len(t2), "S13m_rows": len(t5), "S13h_rows": len(t3)})
    return stats


def readme_rows(x1):
    rows = [["Supplementary tables - Kazakhstan's cybersecurity research, Scopus 2003-2025"],
            ["Date", L.FIXED_DATE],
            ["Corpus", "Scopus, 841 records with a Kazakhstan affiliation published in 2003-2025, retrieved 13 October 2026; Web of Science "
                       "Core Collection as validation and OpenAlex at an equal query (S1)"],
            ["Built by", "phases F13.7a and F13.7c from the tables of closed phases; every number is copied from those tables, none computed here"],
            ["Licence", "Abstracts and the texts of reference lists are not included: the Scopus licence does not permit their redistribution"],
            ["Personal data", "Aggregates and names of institutions, countries, sources, keywords and cited works only; no author of the corpus is "
                              "named (a cited work's first author who is also a corpus author is shown as '" + L.WITHHELD + "')"],
            ["Contents", "; ".join(f"{a} {b}" for a, b in CONTENTS)],
            [],
            ["sheet", "contents", "source (phase / table)", "rows", "sha256 of the source (first 16)"]]
    for r in x1:
        rows.append([r["sheet"], r["description"], r["source_file"] + ("" if r["origin"] == "the file itself" else " <- " + r["origin"]), r["rows"],
                     r["sha256"][:16]])
    return rows, 9


def rewrite_readme(raw, x1):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["README"]
    ws.delete_rows(1, ws.max_row)
    rows, head = readme_rows(x1)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if i in (1, head):
                cell.font = Font(bold=True)
    for col, w in (("A", 32), ("B", 80), ("C", 90), ("D", 8), ("E", 18)):
        ws.column_dimensions[col].width = w
    wb.properties.creator = "phases F13.7a and F13.7c"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    L.install_r17_guard()
    figs = L.figure_order()
    WB.SHEETS = L.sheets_v2(figs)
    WB.DERIVED_ORIGIN = {**WB.DERIVED_ORIGIN, **L.DERIVED_ORIGIN_V2}
    stats = derive_v2()
    raw, x1 = WB.build()
    L.SUPP.mkdir(parents=True, exist_ok=True)
    WB.XLSX.write_bytes(WB.deterministic(rewrite_readme(raw, x1)))
    L.table("X2_sheet_sources", "X2 - sheets of Supplementary_tables.xlsx and their source files", ["sheet", "source_file", "origin", "rows", "columns",
            "sha256", "description"], x1)
    log = {"sheets": len(x1) + 1, "data_sheets": len(x1), "rows_total": sum(r["rows"] for r in x1), "xlsx_bytes": WB.XLSX.stat().st_size,
           "longest_sheet_name": max(len(r["sheet"]) for r in x1), "figures": len(figs), **stats,
           "groups": {g: sum(1 for r in x1 if group(r["sheet"]) == g) for g in [c[0] for c in CONTENTS]}}
    L.jdump(log, L.RUN_LOGS / "workbook.json")
    print(json.dumps({k: log[k] for k in ("sheets", "data_sheets", "rows_total", "xlsx_bytes", "longest_sheet_name", "figures", "S13d_withheld",
                                          "S13m_withheld")}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
