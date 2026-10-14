import sys; sys.dont_write_bytecode = True
"""Independent verifier of phase F13.7c (checks C01-C05).

Written from scratch by the verifier; it imports no phase program and reads nothing in
phase_F13_7c/programs/ except itself. Writes only
outputs/dprg_revision_2026/phase_F13_7c/phase_F13_7c_independent_verification.csv.
R13: corpus Authors are read in memory only; no name, e-mail or ORCID is printed or written;
only counts, sheet names, column headers and row numbers are reported.
"""
import csv
import hashlib
import json
import os
import re
import struct
from collections import Counter, defaultdict

import openpyxl

ROOT = os.environ.get("PROJECT_ROOT", ".")
DPR = 'outputs/dprg_revision_2026/'
PH = DPR + 'phase_F13_7c/'
OUT = os.path.join(ROOT, *(PH + 'phase_F13_7c_independent_verification.csv').split('/'))
csv.field_size_limit(2 ** 31 - 1)

FORBIDDEN = ('outputs/dprg_revision_2026/phase_F13_6d', 'outputs/jcp_submission', 'tools/', '.claude',
             'outputs/dprg_revision_2026/PHASE_REGISTRY.csv', 'main.qmd', 'references.bib',
             'CORRECTIONS.csv', 'PLAN_DPRG_MASTER_2026_09.md', 'HANDOFF_DPRG_2026_09_11.md',
             'PLAN_SJAITU_META_ANALYSIS_PROFESSOR.md', 'PROJECT_STRUCTURE_PROFESSOR.md', '.env',
             'outputs/dprg_revision_2026/phase_F13_7c/programs')


def forbidden(rel):
    r = rel.replace('\\', '/')
    while r.startswith('./'):
        r = r[2:]
    return any(r == f or r.startswith(f.rstrip('/') + '/') or r.endswith('/' + f) for f in FORBIDDEN)


def ap(rel):
    if forbidden(rel):
        raise RuntimeError('refusing a forbidden path')
    return os.path.join(ROOT, *rel.replace('\\', '/').split('/'))


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.reader(f))


RESULTS = []


def rec(check, what, observed, verdict):
    RESULTS.append([check, what, observed, verdict])
    print(f'{check} | {verdict} | {what} | {observed}')


def short(items, n=25):
    items = list(items)
    s = '; '.join(str(x) for x in items[:n])
    return s + (f'; ... (+{len(items) - n})' if len(items) > n else '')


# ---------------------------------------------------------------- cell rule
INT_RE = re.compile(r'^-?(0|[1-9][0-9]*)$')
DEC_RE = re.compile(r'^-?[0-9]+\.([0-9]+)$')
WITHHELD = 'withheld (R13)'


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def data_cell_ok(s, cell):
    """Source string s vs workbook cell under the storage rule. Returns (ok, kind)."""
    v = cell.value
    if cell.data_type == 'f':
        return False, 'formula'
    if s == '':
        return v is None, 'empty'
    if INT_RE.match(s) and str(int(s)) == s:
        return (is_num(v) and float(v).is_integer() and int(v) == int(s)
                and not (isinstance(v, float) and abs(v) > 2 ** 53)), 'int'
    m = DEC_RE.match(s)
    if m:
        n = len(m.group(1))
        ok = is_num(v) and format(float(v), f'.{n}f') == s
        fmt_ok = cell.number_format == '0.' + '0' * n
        if ok and not fmt_ok:
            return False, 'decimal-format'
        return ok, 'decimal'
    return isinstance(v, str) and v == s, 'text'


def header_cell_ok(s, cell):
    v = cell.value
    if s == '':
        return v is None, 'header-empty'
    return isinstance(v, str) and v == s and cell.data_type == 's', 'header'


def compare_sheet(ws, rows):
    """Cell-by-cell comparison of a sheet with CSV rows (row 1 = header)."""
    nrows = len(rows)
    ncols = max((len(r) for r in rows), default=0)
    maxr = max(nrows, ws.max_row)
    maxc = max(ncols, ws.max_column)
    bad = Counter()
    locs = []
    n = 0
    for r_i, wsrow in enumerate(ws.iter_rows(min_row=1, max_row=maxr, min_col=1, max_col=maxc), start=1):
        for c_i, cell in enumerate(wsrow, start=1):
            n += 1
            if r_i <= nrows and c_i <= len(rows[r_i - 1]):
                s = rows[r_i - 1][c_i - 1]
                ok, kind = header_cell_ok(s, cell) if r_i == 1 else data_cell_ok(s, cell)
            else:
                ok, kind = cell.value is None, 'outside-source'
            if not ok:
                bad[kind] += 1
                if len(locs) < 10:
                    locs.append(f'r{r_i}c{c_i}:{kind}')
    dims_ok = (ws.max_row == nrows and ws.max_column == ncols)
    return n, bad, locs, dims_ok, (ws.max_row, ws.max_column), (nrows, ncols)


# ====================================================================== C01
stamp = json.load(open(ap(PH + 'F13_7c_brief_stamp.json'), encoding='utf-8'))
inputs = stamp['inputs']
n_ok = 0
bad = []
skipped = []
for rel, d in inputs.items():
    if forbidden(rel):
        skipped.append(rel)
        continue
    p = ap(rel)
    if not os.path.isfile(p):
        bad.append(rel + ' (missing)')
        continue
    if sha(p) == d['sha256'] and os.path.getsize(p) == d['bytes']:
        n_ok += 1
    else:
        bad.append(rel)
brief_p = ap(PH + 'F13_7c_brief.md')
brief_sha = sha(brief_p)
brief_ok = brief_sha == stamp['brief_sha256'] and os.path.getsize(brief_p) == stamp.get('brief_bytes', os.path.getsize(brief_p))
count_ok = stamp.get('inputs_count', len(inputs)) == len(inputs)
rec('C01', f'sha256+bytes of the {len(inputs)} stamp inputs',
    f'{n_ok} equal; {len(bad)} differ/missing{(" (" + short(bad) + ")") if bad else ""}; '
    f'{len(skipped)} not hashed (forbidden path); inputs_count field {"agrees" if count_ok else "disagrees"}',
    'PASS' if not bad and not skipped and count_ok else 'FAIL')
rec('C01', 'sha256+bytes of F13_7c_brief.md vs stamp brief_sha256/brief_bytes',
    'equal' if brief_ok else 'differ', 'PASS' if brief_ok else 'FAIL')

# ====================================================================== C02
x2 = list(csv.DictReader(open(ap(PH + 'tables/X2_sheet_sources.csv'), encoding='utf-8-sig', newline='')))
wb = openpyxl.load_workbook(ap(PH + 'supplementary/Supplementary_tables.xlsx'))
wb_sheets = [s for s in wb.sheetnames if s != 'README']
x2_sheets = [r['sheet'] for r in x2]

sha_bad, rows_bad, cols_bad, missing = [], [], [], []
src_rows = {}
for r in x2:
    rel = r['source_file']
    p = ap(rel)
    if not os.path.isfile(p):
        missing.append(r['sheet'])
        continue
    if sha(p) != r['sha256']:
        sha_bad.append(r['sheet'])
    rows = read_csv(p)
    src_rows[r['sheet']] = rows
    if len(rows) - 1 != int(r['rows']):
        rows_bad.append(f"{r['sheet']} ({len(rows) - 1} vs {r['rows']})")
    if len(rows[0]) != int(r['columns']):
        cols_bad.append(f"{r['sheet']} ({len(rows[0])} vs {r['columns']})")
rec('C02a', f'sha256 of source_file vs X2 sha256, all {len(x2)} X2 rows',
    f'{len(x2) - len(sha_bad) - len(missing)} equal; {len(sha_bad)} differ{(" (" + short(sha_bad) + ")") if sha_bad else ""}; '
    f'{len(missing)} missing{(" (" + short(missing) + ")") if missing else ""}',
    'PASS' if not sha_bad and not missing else 'FAIL')
rec('C02a', 'data-row count of source_file (lines after header) vs X2 rows',
    f'{len(src_rows) - len(rows_bad)} equal; {len(rows_bad)} differ{(" (" + short(rows_bad) + ")") if rows_bad else ""}',
    'PASS' if not rows_bad and not missing else 'FAIL')
rec('C02a', 'header column count of source_file vs X2 columns (extra)',
    f'{len(src_rows) - len(cols_bad)} equal; {len(cols_bad)} differ{(" (" + short(cols_bad) + ")") if cols_bad else ""}',
    'PASS' if not cols_bad else 'FAIL')
dup_x2 = [s for s, c in Counter(x2_sheets).items() if c > 1]
dup_wb = [s for s, c in Counter(wb.sheetnames).items() if c > 1]
only_wb = sorted(set(wb_sheets) - set(x2_sheets))
only_x2 = sorted(set(x2_sheets) - set(wb_sheets))
rec('C02a', 'workbook sheets (minus README) vs X2 sheets, both directions',
    f'workbook {len(wb.sheetnames)} sheets incl. README ({len(wb_sheets)} data); X2 {len(x2_sheets)}; '
    f'only in workbook {len(only_wb)}{(" (" + short(only_wb) + ")") if only_wb else ""}; only in X2 {len(only_x2)}'
    f'{(" (" + short(only_x2) + ")") if only_x2 else ""}; duplicates X2 {len(dup_x2)}, workbook {len(dup_wb)}; '
    f'same order {"yes" if wb_sheets == x2_sheets else "no"}; README is sheet 1 {"yes" if wb.sheetnames[0] == "README" else "no"}',
    'PASS' if not only_wb and not only_x2 and not dup_x2 and not dup_wb and 'README' in wb.sheetnames else 'FAIL')

derived = [r for r in x2 if r['source_file'].startswith(PH + 'tables/')]
rec('C02a', 'derived sheets (source in phase_F13_7c/tables) carry an origin note',
    f'{len(derived)} derived ({short(r["sheet"] for r in derived)}); with origin text {sum(1 for r in derived if r["origin"].strip())}',
    'INFO')

FIVE = ['S5_inst_thesaurus_v31', 'S5c_inst_v31_changes', 'S6d_collab_by_period_v31',
        'S13j_F133b_T4b_coupling_clust', 'figure_4_collaboration_v31']
x2_by_sheet = {r['sheet']: r for r in x2}
five_fail = 0
for s in FIVE:
    if s not in wb.sheetnames or s not in src_rows:
        rec('C02b', f'{s} cell by cell vs source', 'sheet or source missing', 'FAIL')
        five_fail += 1
        continue
    n, badc, locs, dims_ok, wdim, sdim = compare_sheet(wb[s], src_rows[s])
    ok = not badc and dims_ok
    five_fail += 0 if ok else 1
    rec('C02b', f'{s} cell by cell vs {x2_by_sheet[s]["source_file"].split("/")[-3]}/{x2_by_sheet[s]["source_file"].split("/")[-1]}',
        f'{n} cells compared ({sdim[0] - 1} data rows x {sdim[1]} cols); mismatches {sum(badc.values())}'
        f'{(" " + str(dict(badc)) + " at " + short(locs)) if badc else ""}; sheet dims {wdim} vs source {sdim}',
        'PASS' if ok else 'FAIL')

# extra: every other sheet compared the same way (INFO)
all_ok, all_bad = 0, []
for s in x2_sheets:
    if s in FIVE or s not in wb.sheetnames or s not in src_rows:
        continue
    n, badc, locs, dims_ok, wdim, sdim = compare_sheet(wb[s], src_rows[s])
    if not badc and dims_ok:
        all_ok += 1
    else:
        all_bad.append(f'{s}:{sum(badc.values())}{dict(badc)}{"" if dims_ok else " dims " + str(wdim) + "vs" + str(sdim)}')
rec('C02b+', 'extra: the other X2 sheets cell by cell vs their source_file, same rule',
    f'{all_ok} of {len(x2_sheets) - len(FIVE)} identical; differing {len(all_bad)}{(" (" + short(all_bad) + ")") if all_bad else ""}',
    'INFO')

# C02c: S13d vs the closed phase's T2_top30_works.csv
orig = read_csv(ap(DPR + 'phase_F13_3b/tables/T2_top30_works.csv'))
ws = wb['S13d_F133b_T2_top30_works']
grid = [[c for c in row] for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, len(orig)),
                                                   min_col=1, max_col=max(ws.max_column, len(orig[0])))]
hdr = orig[0]
problems = Counter()
plocs = []
withheld = 0
withheld_not_corpus = 0
corpus_yes_not_withheld = 0
title_cut = 0
title_ws_norm = 0
fa_equal = 0
dims_ok = ws.max_row == len(orig) and ws.max_column == len(hdr)
for c_i, h in enumerate(hdr):
    ok, _ = header_cell_ok(h, grid[0][c_i])
    if not ok:
        problems['header'] += 1
        plocs.append(f'r1c{c_i + 1}')
i_title = hdr.index('title')
i_fa = hdr.index('first_author')
i_cr = hdr.index('corpus_record')
for r_i in range(1, len(orig)):
    o = orig[r_i]
    row = grid[r_i]
    for c_i, s in enumerate(o):
        cell = row[c_i]
        if c_i == i_title:
            words = s.split()
            exp = ' '.join(words[:15]) + ' ...' if len(words) > 15 else s
            v = cell.value
            if len(words) > 15:
                title_cut += 1
            if isinstance(v, str) and v == exp:
                pass
            elif len(words) <= 15 and isinstance(v, str) and v == ' '.join(words):
                title_ws_norm += 1
            elif s == '' and v is None:
                pass
            else:
                problems['title'] += 1
                plocs.append(f'r{r_i + 1}:title')
        elif c_i == i_fa:
            v = cell.value
            if isinstance(v, str) and v == WITHHELD:
                withheld += 1
                if o[i_cr] != 'yes':
                    withheld_not_corpus += 1
            elif (s == '' and v is None) or (isinstance(v, str) and v == s):
                fa_equal += 1
                if o[i_cr] == 'yes':
                    corpus_yes_not_withheld += 1
            else:
                problems['first_author'] += 1
                plocs.append(f'r{r_i + 1}:first_author')
        else:
            ok, kind = data_cell_ok(s, cell)
            if not ok:
                problems[hdr[c_i] + ':' + kind] += 1
                plocs.append(f'r{r_i + 1}:{hdr[c_i]}')
    for c_i in range(len(o), len(row)):
        if row[c_i].value is not None:
            problems['extra-cell'] += 1
n_yes = sum(1 for o in orig[1:] if o[i_cr] == 'yes')
ok_c = not problems and dims_ok and corpus_yes_not_withheld == 0
rec('C02c', 'S13d_F133b_T2_top30_works vs phase_F13_3b/tables/T2_top30_works.csv (title 15-word rule, first_author equal or withheld)',
    f'{len(orig) - 1} rows x {len(hdr)} cols; dims {"equal" if dims_ok else "differ"}; other-cell mismatches {sum(problems.values())}'
    f'{(" " + str(dict(problems)) + " at " + short(plocs)) if problems else ""}; titles over 15 words (cut) {title_cut}; '
    f'titles equal only after whitespace collapse {title_ws_norm}; first_author withheld {withheld}, equal to original {fa_equal}; '
    f'corpus_record=yes rows {n_yes}, of them not withheld {corpus_yes_not_withheld}; withheld with corpus_record=no {withheld_not_corpus}',
    'PASS' if ok_c else 'FAIL')
x2_s13d = x2_by_sheet.get('S13d_F133b_T2_top30_works', {})
rec('C02c', 'X2 origin of S13d names phase_F13_3b T2_top30_works',
    'yes' if 'phase_F13_3b' in x2_s13d.get('origin', '') and 'T2_top30_works' in x2_s13d.get('origin', '') else 'no',
    'INFO')

# ====================================================================== C03
draft_rel = DPR + 'phase_F13_6c/draft/main_v2.qmd'
draft_lines = open(ap(draft_rel), encoding='utf-8').read().splitlines()
FIG_RE = re.compile(r'^!\[(?P<cap>.*)\]\((?P<path>[^)\s]+)\)\{#(?P<lab>fig-[A-Za-z0-9_-]+)[^}]*\}\s*$')
figs = []
for i, line in enumerate(draft_lines, start=1):
    m = FIG_RE.match(line)
    if m:
        figs.append({'line': i, 'cap': m.group('cap'), 'path': m.group('path'), 'label': '@' + m.group('lab')})
near = sum(1 for l in draft_lines if '{#fig-' in l) - len(figs)
nfig = len(figs)


def clean_caption(c):
    return ' '.join(re.sub(r'\[\[.*?\]\]', ' ', c).split())


marks = [len(re.findall(r'\[\[.*?\]\]', f['cap'])) for f in figs]
cap_lines = open(ap(PH + 'figures/figure_captions.txt'), encoding='utf-8-sig').read().splitlines()
cap_hdr = cap_lines[0].split('\t')
cap_rows = [l.split('\t') for l in cap_lines[1:] if l.strip()]
exp_hdr = ['number', 'label', 'caption', 'source_tif', 'sha256_tif', 'source_png', 'sha256_png']
bad_width = sum(1 for r in cap_rows if len(r) != 7)
capd = [dict(zip(cap_hdr, r)) for r in cap_rows]
order_ok = [r.get('label') for r in capd] == [f['label'] for f in figs]
num_ok = [r.get('number') for r in capd] == [f'Figure {k}' for k in range(1, nfig + 1)]
rec('C03', 'figure order: draft image lines (labels) vs figure_captions.txt rows',
    f'draft figures {nfig} (lines {short(f["line"] for f in figs)}; other lines with {{#fig- {near}); captions rows {len(capd)}; '
    f'header as specified {"yes" if cap_hdr == exp_hdr else "no"}; rows with 7 fields {len(capd) - bad_width} of {len(capd)}; '
    f'label order equal {"yes" if order_ok else "no"}; numbers Figure 1..{nfig} {"yes" if num_ok else "no"}',
    'PASS' if order_ok and num_ok and cap_hdr == exp_hdr and bad_width == 0 and nfig > 0 else 'FAIL')
cap_bad = []
cap_alt = []
for k, (f, r) in enumerate(zip(figs, capd), start=1):
    exp = clean_caption(f['cap'])
    if r.get('caption') != exp:
        alt = re.sub(r'\s+([.,;:)])', r'\1', exp)
        (cap_alt if r.get('caption') == alt else cap_bad).append(f'Figure {k}')
rec('C03', 'captions = draft caption minus [[...]] marks, whitespace collapsed',
    f'{nfig - len(cap_bad) - len(cap_alt)} of {nfig} equal; equal only if space before punctuation also dropped {len(cap_alt)}'
    f'{(" (" + short(cap_alt) + ")") if cap_alt else ""}; unequal {len(cap_bad)}{(" (" + short(cap_bad) + ")") if cap_bad else ""}; '
    f'editorial marks per draft caption {marks}',
    'PASS' if not cap_bad and not cap_alt and len(capd) == nfig else 'FAIL')
fig_sheets = [s for s in wb.sheetnames if s.startswith('figure_')]
ks = []
for s in fig_sheets:
    m = re.match(r'^figure_(\d+)_', s)
    ks.append(int(m.group(1)) if m else -1)
sheets_order_ok = ks == sorted(ks) and set(ks) == set(range(1, nfig + 1))
rec('C03', 'workbook sheets named figure_* appear in figure order and cover figures 1..n',
    f'{len(fig_sheets)} figure sheets, figure numbers in workbook order {ks}; non-decreasing {"yes" if ks == sorted(ks) else "no"}; '
    f'cover 1..{nfig} {"yes" if set(ks) == set(range(1, nfig + 1)) else "no"}',
    'PASS' if sheets_order_ok else 'FAIL')
fig_dir = ap(PH + 'figures')
fig_files = sorted(os.listdir(fig_dir))
tifs = sorted(int(m.group(1)) for m in (re.match(r'^figure_(\d+)\.tif$', x) for x in fig_files) if m)
pngs = sorted(int(m.group(1)) for m in (re.match(r'^figure_(\d+)\.png$', x) for x in fig_files) if m)
others = [x for x in fig_files if not re.match(r'^figure_\d+\.(tif|png)$', x) and x != 'figure_captions.txt']
pairs_ok = tifs == pngs == list(range(1, nfig + 1))
rec('C03', 'figure_k.tif/.png pairs in figures/ vs number of draft figures',
    f'tif {tifs}; png {pngs}; draft figures {nfig}; other files besides figure_captions.txt {len(others)}',
    'PASS' if pairs_ok and not others else 'FAIL')

# ====================================================================== C04
f2 = list(csv.DictReader(open(ap(PH + 'tables/F2_figure_set.csv'), encoding='utf-8-sig', newline='')))
f2_by_label = {r['label']: r for r in f2}
draft_dir = os.path.dirname(ap(draft_rel))
COLLAB = '@fig-collaboration-type-by-year'
COLLAB_SRC = DPR + 'phase_F13_7a/figures/Fig_collaboration_type_by_year_v31'


def relof(p):
    return os.path.relpath(p, ROOT).replace(os.sep, '/')


def tif_meta(p):
    with open(p, 'rb') as f:
        b = f.read()
    e = '<' if b[:2] == b'II' else '>'
    off = struct.unpack(e + 'I', b[4:8])[0]
    n = struct.unpack(e + 'H', b[off:off + 2])[0]
    tags = {}
    for i in range(n):
        ent = b[off + 2 + 12 * i: off + 14 + 12 * i]
        tag, typ, cnt = struct.unpack(e + 'HHI', ent[:8])
        if typ == 3:
            val = struct.unpack(e + 'H', ent[8:10])[0]
        elif typ == 4:
            val = struct.unpack(e + 'I', ent[8:12])[0]
        elif typ == 5:
            o = struct.unpack(e + 'I', ent[8:12])[0]
            num, den = struct.unpack(e + 'II', b[o:o + 8])
            val = num / den if den else 0
        else:
            val = None
        tags[tag] = val
    unit = tags.get(296, 2)
    dpi = tags.get(282)
    if dpi is not None and unit == 3:
        dpi *= 2.54
    return tags.get(256), tags.get(257), (round(dpi) if dpi else None)


def png_meta(p):
    with open(p, 'rb') as f:
        b = f.read()
    w, h = struct.unpack('>II', b[16:24])
    dpi = None
    i = 8
    while i < len(b):
        ln = struct.unpack('>I', b[i:i + 4])[0]
        typ = b[i + 4:i + 8]
        if typ == b'pHYs':
            px, py, unit = struct.unpack('>IIB', b[i + 8:i + 17])
            if unit == 1:
                dpi = round(px * 0.0254)
            break
        if typ == b'IDAT':
            break
        i += 12 + ln
    return w, h, dpi


c04_bad = []
f2_bad = []
cap_src_bad = []
meta_notes = []
for k, f in enumerate(figs, start=1):
    if f['label'] == COLLAB:
        stem_rel = COLLAB_SRC
    else:
        resolved = os.path.normpath(os.path.join(draft_dir, f['path'].replace('/', os.sep)))
        stem_rel = os.path.splitext(relof(resolved))[0]
    for ext in ('tif', 'png'):
        src_rel = stem_rel + '.' + ext
        src_p = ap(src_rel)
        dst_p = ap(PH + f'figures/figure_{k}.{ext}')
        if not os.path.isfile(src_p) or not os.path.isfile(dst_p):
            c04_bad.append(f'figure_{k}.{ext} (file missing)')
            continue
        s_src, s_dst = sha(src_p), sha(dst_p)
        if s_src != s_dst:
            c04_bad.append(f'figure_{k}.{ext}')
        r2 = f2_by_label.get(f['label'], {})
        if (r2.get('number') != str(k) or r2.get(f'source_{ext}') != src_rel or r2.get(f'sha256_{ext}') != s_src
                or r2.get(f'bytes_{ext}') != str(os.path.getsize(dst_p))):
            f2_bad.append(f'figure_{k}.{ext}')
        rc = capd[k - 1] if k - 1 < len(capd) else {}
        if rc.get(f'source_{ext}') != src_rel or rc.get(f'sha256_{ext}') != s_src:
            cap_src_bad.append(f'figure_{k}.{ext}')
        w, h, dpi = tif_meta(dst_p) if ext == 'tif' else png_meta(dst_p)
        if (str(w), str(h), str(dpi)) != (r2.get(f'width_px_{ext}'), r2.get(f'height_px_{ext}'), r2.get(f'dpi_{ext}')):
            meta_notes.append(f'figure_{k}.{ext} header {w}x{h}@{dpi} vs F2 {r2.get(f"width_px_{ext}")}x{r2.get(f"height_px_{ext}")}@{r2.get(f"dpi_{ext}")}')
rec('C04', 'sha256 of figures/figure_k.tif/.png = sha256 of its source (draft image path, v3.1 file of F13.7a for the collaboration figure)',
    f'{2 * nfig - len(c04_bad)} of {2 * nfig} files equal{(" ; differ: " + short(c04_bad)) if c04_bad else ""}',
    'PASS' if not c04_bad and nfig > 0 else 'FAIL')
rec('C04', 'F2_figure_set.csv: number, source_tif/png, sha256_tif/png, bytes_tif/png agree with the derived sources and copies',
    f'{len(f2)} F2 rows; disagreeing {len(f2_bad)}{(" (" + short(f2_bad) + ")") if f2_bad else ""}; labels in draft order '
    f'{"yes" if [r["label"] for r in f2] == [f["label"] for f in figs] else "no"}',
    'PASS' if not f2_bad and len(f2) == nfig else 'FAIL')
rec('C04', 'figure_captions.txt source_tif/png and sha256_tif/png agree with the derived sources',
    f'disagreeing {len(cap_src_bad)}{(" (" + short(cap_src_bad) + ")") if cap_src_bad else ""}',
    'PASS' if not cap_src_bad else 'FAIL')
rec('C04', 'extra: width/height/dpi read from the copies\' file headers (no image opened) vs F2',
    f'{2 * nfig - len(meta_notes)} of {2 * nfig} agree{(" ; " + short(meta_notes)) if meta_notes else ""}', 'INFO')
collab = [f for f in figs if f['label'] == COLLAB]
if collab:
    dp = os.path.normpath(os.path.join(draft_dir, collab[0]['path'].replace('/', os.sep)))
    ex = os.path.isfile(dp)
    same = ex and sha(dp) == sha(ap(COLLAB_SRC + '.tif'))
    rec('C04', 'collaboration figure: draft path (pre-v3.1 file) vs the v3.1 file used',
        f'draft path {relof(dp)} exists {"yes" if ex else "no"}; byte-identical to the v3.1 tif {"yes" if same else "no"} '
        f'(the brief prescribes the v3.1 file)', 'INFO')
f2_sheets_ok = all(set(x.strip() for x in r['data_sheets'].split(';')) == {s for s in fig_sheets if s.startswith(f'figure_{r["number"]}_')}
                   for r in f2)
rec('C04', 'F2 data_sheets column vs workbook figure_k_* sheets', 'agree' if f2_sheets_ok else 'disagree', 'INFO')

# ====================================================================== C05
INITIAL = re.compile(r'[^\W\d_]+\.')
LETTERS = re.compile(r'[^\W\d_]+')
surnames = set()
fullnames = set()
n_names = 0
with open(ap(DPR + 'author_exports_2026_10/scopus_full_TAK_KZ_2003_2025_2026-10-13.csv'), encoding='utf-8-sig', newline='') as f:
    for row in csv.DictReader(f):
        for nm in row['Authors'].split(';'):
            nm = nm.strip()
            if not nm:
                continue
            n_names += 1
            fullnames.add(nm)
            for t in LETTERS.findall(INITIAL.sub(' ', nm)):
                if len(t) >= 4 and t[0].isupper():
                    surnames.add(t)
ORG_SOURCES = [(DPR + 'phase_F13_2e/tables/T3a_v31_kz_institutions.csv', ['institution']),
               (DPR + 'phase_F13_2b/tables/T3e_foreign_institutions.csv', ['institution']),
               (DPR + 'phase_F13_2b/tables/scopus_country_names.csv', ['country_name']),
               (DPR + 'phase_F13_2/tables/T2a_sources.csv', ['source_title', 'publisher']),
               (DPR + 'phase_F13_2/tables/T2d_publishers.csv', ['publisher'])]
org_words = set()
missing_cols = []
for rel, cols in ORG_SOURCES:
    with open(ap(rel), encoding='utf-8-sig', newline='') as f:
        rd = csv.DictReader(f)
        for c in cols:
            if c not in rd.fieldnames:
                missing_cols.append(rel.split('/')[-1] + ':' + c)
        for row in rd:
            for c in cols:
                for w in LETTERS.findall(row.get(c) or ''):
                    org_words.add(w.lower())
set_aside = {t for t in surnames if t.lower() in org_words}
tokens = surnames - set_aside
full_by_token = defaultdict(set)
for nm in fullnames:
    for t in LETTERS.findall(INITIAL.sub(' ', nm)):
        if t in tokens:
            full_by_token[t].add(nm)
rec('C05', 'corpus surname-token set (Authors, letters only, len>=4, capitalised, initials stripped) minus org/place words',
    f'{n_names} author entries, {len(fullnames)} distinct; {len(surnames)} tokens; set aside {len(set_aside)} '
    f'(org/place words from 5 tables, {len(org_words)} distinct words); kept {len(tokens)}; '
    f'missing columns {len(missing_cols)}{(" (" + short(missing_cols) + ")") if missing_cols else ""}',
    'INFO' if not missing_cols else 'FAIL')

WORD = re.compile(r'\w+')
EMAIL = re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)+')
ORCID = re.compile(r'\b\d{4}-\d{4}-\d{4}-\d{3}[\dX]\b')
BADCOLS = {'authors', 'author', 'email', 'orcid', 'abstract'}

# lowercase vocabulary of the workbook (to say whether a hit token is also used as an ordinary word)
lower_vocab = set()
for ws in wb.worksheets:
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str):
                lower_vocab.update(t for t in WORD.findall(v) if t.islower())

wb_hits = defaultdict(list)       # sheet -> [(row, header, token)]
fullname_hits = []
email_hits, orcid_hits, long_cells, bad_cols = [], [], [], []
withheld_by_sheet = Counter()
named_cells_by_sheet = Counter()
for ws in wb.worksheets:
    headers = {}
    for r_i, row in enumerate(ws.iter_rows(), start=1):
        for c_i, cell in enumerate(row, start=1):
            v = cell.value
            if r_i == 1 and ws.title != 'README':
                headers[c_i] = str(v) if v is not None else ''
                if isinstance(v, str) and v.strip().lower() in BADCOLS:
                    bad_cols.append(f'{ws.title}:{v}')
            if not isinstance(v, str):
                continue
            h = headers.get(c_i, '') if ws.title != 'README' else ''
            if r_i > 1 and h in ('first_author', 'top_work_first_author'):
                if v == WITHHELD:
                    withheld_by_sheet[ws.title] += 1
                else:
                    named_cells_by_sheet[ws.title] += 1
            for t in WORD.findall(v):
                if t in tokens:
                    wb_hits[ws.title].append((r_i, h, t))
                    if any(fn in v for fn in full_by_token.get(t, ())):
                        fullname_hits.append(f'{ws.title}:r{r_i}')
            if EMAIL.search(v):
                email_hits.append(f'{ws.title}:r{r_i}')
            if ORCID.search(v):
                orcid_hits.append(f'{ws.title}:r{r_i}')
            if ws.title != 'README' and len(v.split()) > 30:
                long_cells.append(f'{ws.title}:r{r_i}c{c_i}')

file_hits = defaultdict(list)
files_scanned = []
tab_dir = ap(PH + 'tables')
scan_files = [(PH + 'tables/' + x) for x in sorted(os.listdir(tab_dir)) if os.path.isfile(os.path.join(tab_dir, x))]
scan_files.append(PH + 'figures/figure_captions.txt')
for rel in scan_files:
    files_scanned.append(rel.split('/')[-1])
    text = open(ap(rel), encoding='utf-8', errors='replace').read()
    for l_i, line in enumerate(text.splitlines(), start=1):
        for t in WORD.findall(line):
            if t in tokens:
                file_hits[rel.split('/')[-1]].append((l_i, t))
                if any(fn in line for fn in full_by_token.get(t, ())):
                    fullname_hits.append(f'{rel.split("/")[-1]}:line{l_i}')
        if EMAIL.search(line):
            email_hits.append(f'{rel.split("/")[-1]}:line{l_i}')
        if ORCID.search(line):
            orcid_hits.append(f'{rel.split("/")[-1]}:line{l_i}')
    if rel.endswith('.csv'):
        hdr_f = read_csv(ap(rel))[0] if text.strip() else []
        for hname in hdr_f:
            if hname.strip().lower() in BADCOLS:
                bad_cols.append(f'{rel.split("/")[-1]}:{hname}')

n_wb_hits = sum(len(v) for v in wb_hits.values())
n_file_hits = sum(len(v) for v in file_hits.values())
hit_tokens = {t for v in wb_hits.values() for (_, _, t) in v} | {t for v in file_hits.values() for (_, t) in v}
common = {t for t in hit_tokens if t.lower() in lower_vocab}
in_author_field = sum(1 for v in wb_hits.values() for (_, h, _) in v if h in ('first_author', 'top_work_first_author'))
rec('C05', 'whole-word case-sensitive surname-token hits in every workbook cell (README included)',
    f'{n_wb_hits} hits in {len(wb_hits)} sheets; distinct tokens {len({t for v in wb_hits.values() for (_, _, t) in v})}; '
    f'hits in first_author/top_work_first_author cells {in_author_field}',
    'PASS' if n_wb_hits == 0 else 'INFO')
for s in [x for x in wb.sheetnames if x in wb_hits]:
    hs = wb_hits[s]
    by_col = Counter(h if h else '(README)' for (_, h, _) in hs)
    rows_ = sorted({r for (r, _, _) in hs})
    rec('C05', f'hit locations in sheet {s}',
        f'{len(hs)} hits, {len({t for (_, _, t) in hs})} distinct tokens, of which also used as a lower-case word in the workbook '
        f'{len({t for (_, _, t) in hs} & common)}; by column {dict(by_col)}; rows {short(rows_, 40)}', 'INFO')
rec('C05', f'whole-word surname-token hits in the {len(files_scanned)} files of tables/ and figures/figure_captions.txt',
    f'{n_file_hits} hits in {len(file_hits)} files'
    f'{(" (" + short(f"{k}: {len(v)} hits, lines {short(sorted({l for (l, _) in v}), 20)}" for k, v in file_hits.items()) + ")") if file_hits else ""}',
    'PASS' if n_file_hits == 0 else 'INFO')
rec('C05', 'full corpus author strings ("Surname I.") found verbatim in any cell or file line',
    f'{len(fullname_hits)}{(" (" + short(fullname_hits) + ")") if fullname_hits else ""}',
    'PASS' if not fullname_hits else 'FAIL')
rec('C05', 'e-mail-shaped / ORCID-shaped strings (workbook, tables/, captions)',
    f'e-mail {len(email_hits)}{(" (" + short(email_hits) + ")") if email_hits else ""}; ORCID {len(orcid_hits)}'
    f'{(" (" + short(orcid_hits) + ")") if orcid_hits else ""}',
    'PASS' if not email_hits and not orcid_hits else 'FAIL')
rec('C05', 'columns named authors/author/email/orcid/abstract (sheet headers and tables/*.csv headers)',
    f'{len(bad_cols)}{(" (" + short(bad_cols) + ")") if bad_cols else ""}', 'PASS' if not bad_cols else 'FAIL')
rec('C05', 'cells over 30 words outside README',
    f'{len(long_cells)}{(" (" + short(long_cells) + ")") if long_cells else ""}', 'PASS' if not long_cells else 'FAIL')
# positive control of the token scan: the closed phase's original first_author values of the rows that the
# workbook withholds should mostly hit the token set (counts only; values stay in memory)
ctrl_rows = 0
ctrl_hit = 0
for r_i in range(1, len(orig)):
    if isinstance(grid[r_i][i_fa].value, str) and grid[r_i][i_fa].value == WITHHELD:
        ctrl_rows += 1
        if any(t in tokens for t in WORD.findall(orig[r_i][i_fa])):
            ctrl_hit += 1
rec('C05', 'positive control: the scan applied to the ORIGINAL first_author of the S13d rows the workbook withholds',
    f'{ctrl_hit} of {ctrl_rows} original values hit the token set (scan is live)', 'INFO')
rec('C05', 'first_author / top_work_first_author cells "withheld (R13)" (workbook)',
    f'withheld {sum(withheld_by_sheet.values())} {dict(withheld_by_sheet)}; not withheld (non-empty) '
    f'{sum(named_cells_by_sheet.values())} {dict(named_cells_by_sheet)}', 'INFO')

# ---------------------------------------------------------------- output
with open(OUT, 'w', encoding='utf-8', newline='') as f:
    w = csv.writer(f)
    w.writerow(['check', 'what', 'observed', 'verdict'])
    w.writerows(RESULTS)
print('written', relof(OUT), len(RESULTS), 'rows;', Counter(r[3] for r in RESULTS))
