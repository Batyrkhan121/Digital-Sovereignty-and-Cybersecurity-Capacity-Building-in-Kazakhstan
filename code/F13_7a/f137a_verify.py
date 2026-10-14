"""Independent verification of phase F13.7a (package tools).

Written by the independent verifier; it reads only the phase outputs and their
named input tables, never the phase's own programs. It writes exactly one file:
OUT/phase_F13_7a_independent_verification.csv.

Figures are read for METADATA ONLY (size, dpi, format, PDF MediaBox) - R18.
"""
import sys
sys.dont_write_bytecode = True

import csv
import hashlib
import io
import json
import random
import re
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from PIL import Image

import os
ROOT = Path(os.environ.get("PROJECT_ROOT", "."))
REV = ROOT / "outputs" / "dprg_revision_2026"
OUT = REV / "phase_F13_7a"
RESULT = OUT / "phase_F13_7a_independent_verification.csv"

FORBIDDEN_PARTS = ("phase_F13_3b", "phase_F13_6c", ".claude", "PHASE_REGISTRY.csv",
                   "CORRECTIONS.csv", "PLAN_DPRG_MASTER_2026_09.md",
                   "HANDOFF_DPRG_2026_09_11.md", "PLAN_SJAITU_META_ANALYSIS_PROFESSOR.md",
                   "PROJECT_STRUCTURE_PROFESSOR.md", ".env")
E1_EXPORT = re.compile(r"author_exports_2026_10[/\\]scopus_full_[^/\\]*\.csv$", re.I)


def guard(path):
    """Refuse to open any path the brief forbids."""
    s = str(path).replace("\\", "/")
    parts = s.split("/")
    for f in FORBIDDEN_PARTS:
        if f in parts:
            raise SystemExit(f"refused forbidden path: {s}")
    if E1_EXPORT.search(s):
        raise SystemExit(f"refused forbidden path: {s}")
    if "phase_F13_7a/programs/" in s and not s.endswith("f137a_verify.py"):
        raise SystemExit(f"refused phase program: {s}")
    return path


def sha256(path):
    h = hashlib.sha256()
    with open(guard(path), "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_csv(path):
    with open(guard(path), encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f) if r != []]
    return rows[0], rows[1:]


def read_dicts(path):
    with open(guard(path), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


RESULTS = []


def words(s):
    return len(str(s).split())


def add(check_id, what, expected, recomputed, ok, note=""):
    row = [check_id, what, str(expected), str(recomputed), "PASS" if ok else "FAIL", note]
    for cell in row:
        if words(cell) > 30:
            raise SystemExit(f"cell over 30 words in {check_id}: {cell!r}")
    RESULTS.append(row)


def short(items, n=4):
    items = list(items)
    if not items:
        return "none"
    s = "; ".join(str(i) for i in items[:n])
    if len(items) > n:
        s += f"; +{len(items) - n} more"
    return s


# ---------------------------------------------------------------- C01 digests
stamp = json.loads(guard(OUT / "F13_7a_brief_stamp.json").read_text(encoding="utf-8"))
inputs = stamp["inputs"]
bad = []
for rel, rec in inputs.items():
    p = ROOT / rel
    exp_sha = rec["sha256"] if isinstance(rec, dict) else rec
    exp_bytes = rec.get("bytes") if isinstance(rec, dict) else None
    if not p.is_file():
        bad.append(f"{rel}: missing")
        continue
    if sha256(p) != exp_sha:
        bad.append(f"{rel}: sha256 differs")
    if exp_bytes is not None and p.stat().st_size != exp_bytes:
        bad.append(f"{rel}: bytes differ")
add("C01", "stamp inputs: sha256 (and bytes) equal the files now",
    f"{len(inputs)} inputs as stamped", f"{len(inputs) - len(bad)} of {len(inputs)} equal",
    not bad, short(bad, 2) if bad else "all digests and sizes equal")

# ---------------------------------------------------------------- C02 workbook
x1 = read_dicts(OUT / "tables" / "X1_sheet_sources.csv")
wb = openpyxl.load_workbook(guard(OUT / "supplementary" / "Supplementary_tables_draft.xlsx"))
names = wb.sheetnames

src_bad, col_bad = [], []
for r in x1:
    p = ROOT / r["source_file"]
    if not p.is_file():
        src_bad.append(f"{r['sheet']}: source missing")
        continue
    head, data = read_csv(p)
    if sha256(p) != r["sha256"]:
        src_bad.append(f"{r['sheet']}: sha256")
    if len(data) != int(r["rows"]):
        src_bad.append(f"{r['sheet']}: rows {len(data)} vs {r['rows']}")
    if r.get("columns") and len(head) != int(r["columns"]):
        col_bad.append(r["sheet"])
add("C02a", "X1 source files: sha256 and data-row count equal X1",
    f"{len(x1)} source rows in X1", f"{len(x1) - len(src_bad)} equal",
    not src_bad,
    (short(src_bad, 3) if src_bad else "all equal") +
    (f"; column counts differ: {short(col_bad, 3)}" if col_bad else "; column counts also equal"))

long_names = [n for n in names if len(n) > 31]
x1_sheets = [r["sheet"] for r in x1]
missing_sheets = [s for s in x1_sheets if s not in names]
extra_sheets = [n for n in names if n not in x1_sheets and n != "README"]
ok = (not long_names) and ("README" in names) and not missing_sheets
add("C02b", "sheet names <= 31 characters; sheet README exists; every X1 sheet present",
    f"README + {len(x1_sheets)} X1 sheets", f"{len(names)} sheets, longest {max(len(n) for n in names)} chars",
    ok, f"over 31: {short(long_names)}; missing: {short(missing_sheets)}; not in X1 besides README: {short(extra_sheets)}")

# row/column extent of every X1 sheet against X1
extent_bad = []
for r in x1:
    if r["sheet"] not in names:
        continue
    ws = wb[r["sheet"]]
    if ws.max_row - 1 != int(r["rows"]) or ws.max_column != int(r["columns"]):
        extent_bad.append(f"{r['sheet']} {ws.max_row - 1}x{ws.max_column}")
add("C02c", "every sheet: data rows and columns equal X1 rows and columns",
    f"{len(x1)} sheets", f"{len(x1) - len(extent_bad)} equal", not extent_bad, short(extent_bad, 3))

INT_RE = re.compile(r"^-?(0|[1-9]\d*)$")
DEC_RE = re.compile(r"^-?\d+\.(\d+)$")


def cell_ok(s, cell):
    """Return None when the cell holds the source string s by the brief's rule, else a reason."""
    v = cell.value
    if s == "":
        return None if v in (None, "") else "empty source, cell filled"
    if cell.data_type == "f":
        return "formula cell"
    if INT_RE.match(s):
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            return "integer stored as text"
        if isinstance(v, float) and not v.is_integer():
            return "integer value differs"
        return None if int(v) == int(s) else "integer value differs"
    m = DEC_RE.match(s)
    if m:
        d = len(m.group(1))
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            return "decimal stored as text"
        if format(float(v), f".{d}f") != s:
            return "decimal does not format back"
        fmt = "0." + "0" * d
        if cell.number_format != fmt:
            return f"display format {cell.number_format} not {fmt}"
        return None
    if not isinstance(v, str):
        return "text stored as non-text"
    return None if v == s else "text differs"


fixed = ["S6a_institutions_v31", "S4_kw_thesaurus_v4"]
pool = [s for s in x1_sheets if s not in fixed]
picked = random.Random(0).sample(pool, 5)
src_of = {r["sheet"]: r["source_file"] for r in x1}
for sheet in picked + fixed:
    head, data = read_csv(ROOT / src_of[sheet])
    ws = wb[sheet]
    probs = []
    ncells = 0
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=j)
        ncells += 1
        if not (isinstance(c.value, str) and c.value == h) and not (h == "" and c.value in (None, "")):
            probs.append(f"header col {j}")
    reasons = Counter()
    kinds_seen = Counter()
    for i, row in enumerate(data, start=2):
        for j in range(1, len(head) + 1):
            s = row[j - 1] if j - 1 < len(row) else ""
            ncells += 1
            kinds_seen["empty" if s == "" else "integer" if INT_RE.match(s)
                       else "decimal" if DEC_RE.match(s) else "text"] += 1
            why = cell_ok(s, ws.cell(row=i, column=j))
            if why:
                reasons[why] += 1
                if len(probs) < 3:
                    probs.append(f"r{i}c{j} {why}")
    # nothing beyond the source block
    beyond = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            if (c.row > len(data) + 1 or c.column > len(head)) and c.value not in (None, ""):
                beyond += 1
    if beyond:
        probs.append(f"{beyond} cells outside the source block")
    fails = sum(reasons.values()) + beyond + sum(1 for p in probs if p.startswith("header"))
    tag = "random seed 0" if sheet in picked else "named in brief"
    add("C02d", f"sheet {sheet} ({tag}): every cell equals source CSV by the storage rule",
        f"{len(data)} data rows x {len(head)} columns", f"{ncells} cells compared, {fails} mismatches",
        fails == 0, (short(probs, 3) if probs else "all equal") + "; data cells: " +
        ", ".join(f"{k} {kinds_seen.get(k, 0)}" for k in ("integer", "decimal", "text", "empty")))

# ---------------------------------------------------------------- C03 figure manifest
f1 = read_dicts(OUT / "tables" / "F1_figure_manifest.csv")
MEDIABOX = re.compile(rb"/MediaBox\s*\[\s*([-\d.]+)\s+([-\d.]+)\s+([-\d.]+)\s+([-\d.]+)\s*\]")
n_exist = 0
for r in f1:
    rel = r["file"]
    p = ROOT / rel
    name = Path(rel).name
    if r["exists"] != "1":
        absent = not p.exists()
        add("C03", f"{name}: manifest exists=0 means the file is absent",
            f"exists={r['exists']}, verdict '{r['verdict']}'", f"file present: {not absent}",
            absent and r["verdict"] == "file missing", "row with exists=0; only absence checked")
        continue
    n_exist += 1
    probs = []
    if not p.is_file():
        add("C03", f"{name}: metadata as manifest", "exists=1", "file missing", False, "")
        continue
    size = p.stat().st_size
    if str(size) != r["bytes"]:
        probs.append(f"bytes {size} vs {r['bytes']}")
    fmt = r["format"].lower()
    if fmt == "pdf":
        raw = guard(p).read_bytes()
        if not raw.startswith(b"%PDF"):
            probs.append("no %PDF header")
        m = MEDIABOX.search(raw)
        if not m:
            probs.append("MediaBox not found")
            rec = "MediaBox not found"
        else:
            a, b, c, d = (float(x) for x in m.groups())
            w, h = c - a, d - b
            rec = f"{w:.1f} x {h:.1f} pt, {size} bytes"
            try:
                mw, mh = (float(x) for x in r["page_size_pt"].split(" x "))
                if abs(mw - w) > 0.05 or abs(mh - h) > 0.05:
                    probs.append(f"page size {w} x {h} vs {r['page_size_pt']}")
            except ValueError:
                probs.append(f"page size unreadable: {r['page_size_pt']}")
        exp_verdict = "meets"
        if r["dpi"] != "vector":
            probs.append(f"dpi field {r['dpi']} not 'vector'")
        exp = f"{r['page_size_pt']} pt, {r['bytes']} bytes, verdict {r['verdict']}"
    else:
        with Image.open(guard(p)) as im:
            iw, ih = im.size
            ifmt = (im.format or "").lower()
            dpi = im.info.get("dpi")
        if ifmt != fmt:
            probs.append(f"format {ifmt} vs {fmt}")
        dx = round(float(dpi[0])) if dpi else None
        dy = round(float(dpi[1])) if dpi else None
        if dx != dy:
            probs.append(f"dpi x {dx} differs from y {dy}")
        if str(iw) != r["width_px"] or str(ih) != r["height_px"]:
            probs.append(f"size {iw}x{ih} vs {r['width_px']}x{r['height_px']}")
        if str(dx) != r["dpi"]:
            probs.append(f"dpi {dx} vs {r['dpi']}")
        meets = dx is not None and dx >= 300 and iw >= 1000
        exp_verdict = "meets" if meets else "not meets"
        if r["meets_300_dpi"] != ("yes" if (dx or 0) >= 300 else "no"):
            probs.append("meets_300_dpi field")
        if r["meets_width_1000_px"] != ("yes" if iw >= 1000 else "no"):
            probs.append("meets_width_1000_px field")
        rec = f"{iw}x{ih} px, {dx} dpi, {size} bytes"
        exp = f"{r['width_px']}x{r['height_px']} px, {r['dpi']} dpi, {r['bytes']} bytes, verdict {r['verdict']}"
    if (exp_verdict == "meets") != (r["verdict"] == "meets"):
        probs.append(f"verdict {r['verdict']} vs rule {exp_verdict}")
    add("C03", f"{name}: bytes, size, dpi or MediaBox and verdict rule",
        exp, rec + f", rule verdict {exp_verdict}", not probs, short(probs, 3) if probs else "metadata only; all fields equal")

# ---------------------------------------------------------------- C04 redrawn figure
cols = ["year", "single institution", "domestic multi-institution", "international", "total"]
fig_rows = read_dicts(OUT / "figures" / "Fig_collaboration_type_by_year_v31.csv")
t2f_rows = read_dicts(REV / "phase_F13_2e" / "tables" / "T2f_v31_collaboration_by_year.csv")
diff = []
if len(fig_rows) != len(t2f_rows):
    diff.append(f"rows {len(fig_rows)} vs {len(t2f_rows)}")
for a, b in zip(fig_rows, t2f_rows):
    for c in cols:
        if a.get(c) != b.get(c):
            diff.append(f"year {b.get('year')} {c}")
add("C04a", "figure CSV equals T2f_v31 in year and the four count columns",
    f"{len(t2f_rows)} rows x {len(cols)} columns from T2f_v31", f"{len(fig_rows)} rows, {len(diff)} differences",
    not diff, short(diff, 3) if diff else "all values equal, same order")

stem = OUT / "figures" / "Fig_collaboration_type_by_year_v31"
sizes = {ext: (stem.with_suffix(ext).stat().st_size if stem.with_suffix(ext).is_file() else 0)
         for ext in (".pdf", ".png", ".tif")}
add("C04b", "redrawn figure pdf, png and tif exist with size > 0",
    "three files, each > 0 bytes", ", ".join(f"{k} {v}" for k, v in sizes.items()),
    all(v > 0 for v in sizes.values()), "existence and size only")
dpis = {}
for ext in (".png", ".tif"):
    with Image.open(guard(stem.with_suffix(ext))) as im:
        d = im.info.get("dpi")
    dpis[ext] = tuple(round(float(x)) for x in d) if d else None
add("C04c", "redrawn png and tif carry 300 dpi",
    "300 x 300 dpi each", ", ".join(f"{k} {v}" for k, v in dpis.items()),
    all(v == (300, 300) for v in dpis.values()), "metadata only (R18)")

# ---------------------------------------------------------------- C05 corrections draft
c1_head, _ = read_csv(OUT / "tables" / "C1_corrections_draft.csv")
c1 = read_dicts(OUT / "tables" / "C1_corrections_draft.csv")
n1 = read_dicts(REV / "phase_F13_6a" / "tables" / "N1_manuscript_numbers.csv")
n1_by_id = {r["id"]: r for r in n1}
n1_target = [r for r in n1 if r["status"] in ("changed", "remove")]
ids_c1 = [r["n1_id"] for r in c1]
dups = [k for k, v in Counter(ids_c1).items() if v > 1]
missing = sorted(set(r["id"] for r in n1_target) - set(ids_c1))
extra = sorted(set(ids_c1) - set(r["id"] for r in n1_target))
field_bad = []
for r in c1:
    nr = n1_by_id.get(r["n1_id"])
    if nr is None:
        continue
    if r["old_value"] != nr["number_as_written"]:
        field_bad.append(f"{r['n1_id']} old_value")
    if r["n1_status"] != nr["status"]:
        field_bad.append(f"{r['n1_id']} n1_status")
ok = len(c1) == 118 and len(n1_target) == 118 and not dups and not missing and not extra and not field_bad
add("C05a", "C1 has one row per N1 row of status changed or remove; old_value and n1_status from N1",
    f"118 rows; N1 changed/remove rows {len(n1_target)}", f"{len(c1)} rows, duplicates {len(dups)}, missing {len(missing)}, extra {len(extra)}",
    ok, f"field mismatches: {short(field_bad, 3)}")

want_tail = ["correction_id", "date", "superseded_file", "superseded_field",
             "superseded_value", "corrected_value", "evidence_file", "phase"]
add("C05b", "last eight C1 columns carry the CORRECTIONS format names",
    "; ".join(want_tail), "; ".join(c1_head[-8:]), c1_head[-8:] == want_tail, "names and order compared")

e1 = read_dicts(REV / "phase_F13_6b" / "tables" / "E1_evidence_index.csv")
e1_by_id = {r["id"]: r for r in e1}
v31_cache = {}


def nocomma(s):
    return re.sub(r"(?<=\d),(?=\d{3}(?!\d))", "", str(s).strip())


def v31_value(row):
    """Value named by note_for_F13_6d: the F13.2e table, key and column from source_of_new."""
    m = re.search(r"F13\.2e\s+(\w+)\s+gives", row["note_for_F13_6d"])
    k = re.search(r"\[(.*)\]\s*(.+?)\s*$", row["source_of_new"])
    if not m or not k:
        return None, "note or source not parsable"
    table = m.group(1)
    path = REV / "phase_F13_2e" / "tables" / f"{table}.csv"
    if not path.is_file():
        return None, f"{table} not found"
    if table not in v31_cache:
        v31_cache[table] = read_dicts(path)
    pairs = re.split(r",\s*(?=[A-Za-z_][\w ]*=)", k.group(1))
    keys = dict(p.split("=", 1) for p in pairs)
    col = k.group(2)
    hits = [r for r in v31_cache[table] if all(r.get(a) == b for a, b in keys.items())]
    if len(hits) != 1 or col not in hits[0]:
        return None, f"{len(hits)} key matches in {table}"
    return hits[0][col], table


def check_c1_row(r):
    probs = []
    nr = n1_by_id.get(r["n1_id"])
    if nr is None:
        return ["n1_id not in N1"]
    eids = re.findall(r"E-\d+", r["in_draft_E1"])
    if r["n1_status"] == "remove":
        if r["new_value"] != "removed":
            probs.append(f"{r['n1_id']} removed but new_value '{r['new_value']}'")
        if eids:
            probs.append(f"{r['n1_id']} removed row lists E-ids")
        return probs
    if r["n1_value_new"] != nr["value_new"]:
        probs.append(f"{r['n1_id']} n1_value_new differs from N1")
    if r["new_value"] != r["n1_value_new"]:
        if r["note_for_F13_6d"].strip():
            val, where = v31_value(r)
            if val is None or val != r["new_value"]:
                probs.append(f"{r['n1_id']} new_value not in F13.2e table ({where})")
        else:
            probs.append(f"{r['n1_id']} new_value differs, no note")
    for e in eids:
        er = e1_by_id.get(e)
        if er is None:
            probs.append(f"{r['n1_id']} {e} not in E1")
        elif nocomma(er["value_in_source"]) != nocomma(r["new_value"]):
            probs.append(f"{r['n1_id']} {e} value differs")
    return probs


sample = random.Random(0).sample(c1, 30)
probs = []
for r in sample:
    probs += check_c1_row(r)
st = Counter(r["n1_status"] for r in sample)
n_e = sum(len(re.findall(r"E-\d+", r["in_draft_E1"])) for r in sample)
add("C05c", "sample of 30 C1 rows (seed 0): new_value rule and E1 evidence values",
    "removed rows 'removed'; changed rows from N1 value_new or named F13.2e table; E-ids exist and match",
    f"{st.get('remove', 0)} remove, {st.get('changed', 0)} changed, {n_e} E-ids checked, {len(probs)} problems",
    not probs, short(probs, 3) if probs else "all sampled rows follow the rule")

probs_all = []
for r in c1:
    probs_all += check_c1_row(r)
n_note = sum(1 for r in c1 if r["note_for_F13_6d"].strip())
n_e_all = sum(len(re.findall(r"E-\d+", r["in_draft_E1"])) for r in c1)
add("C05d", "same rule applied to all 118 C1 rows (beyond the sample)",
    "no row breaks the rule", f"{len(c1)} rows, {n_note} with note, {n_e_all} E-ids, {len(probs_all)} problems",
    not probs_all, short(probs_all, 3) if probs_all else "all rows follow the rule; notes resolved in the F13.2e table")

tail_bad = []
for r in c1:
    if r["superseded_value"] != r["old_value"]:
        tail_bad.append(f"{r['n1_id']} superseded_value")
    if r["n1_status"] == "changed" and r["corrected_value"] != r["new_value"]:
        tail_bad.append(f"{r['n1_id']} corrected_value")
    if r["n1_status"] == "remove" and not r["corrected_value"].startswith("removed"):
        tail_bad.append(f"{r['n1_id']} corrected_value")
add("C05e", "CORRECTIONS-format fields agree with old_value and new_value",
    "superseded_value = old_value; corrected_value = new_value (removed rows start 'removed')",
    f"{len(tail_bad)} disagreements in {len(c1)} rows", not tail_bad, short(tail_bad, 3))

# ---------------------------------------------------------------- C06 R13
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
ORCID = re.compile(r"\b(?:0000|0009)-\d{4}-\d{4}-\d{3}[\dX]\b")
files = sorted(p for p in (OUT / "tables").iterdir() if p.is_file())
hits = []
for p in files:
    t = guard(p).read_text(encoding="utf-8", errors="replace")
    if EMAIL.search(t):
        hits.append(f"{p.name} e-mail shape")
    if ORCID.search(t):
        hits.append(f"{p.name} ORCID shape")
add("C06a", "no e-mail or ORCID shapes in OUT tables (all files)",
    "0 hits", f"{len(files)} files scanned, {len(hits)} hits", not hits, short(hits, 3))

figcsv = sorted(p for p in (OUT / "figures").iterdir() if p.is_file() and p.suffix.lower() == ".csv")
hits = []
for p in figcsv:
    t = guard(p).read_text(encoding="utf-8", errors="replace")
    if EMAIL.search(t) or ORCID.search(t):
        hits.append(p.name)
add("C06b", "no e-mail or ORCID shapes in OUT figures CSV files",
    "0 hits", f"{len(figcsv)} files scanned, {len(hits)} hits", not hits, short(hits, 3))

hits, long_cells, ncells = [], [], 0
banned_hdr, author_hdr = [], set()
BANNED = {"authors", "author full names", "abstract", "email", "e-mail", "orcid"}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            ncells += 1
            s = str(c.value)
            if EMAIL.search(s) or ORCID.search(s):
                hits.append(f"{ws.title}!{c.coordinate}")
            if ws.title != "README" and words(s) > 30:
                long_cells.append(f"{ws.title}!{c.coordinate}")
    if ws.title == "README":
        continue
    for c in ws[1]:
        if isinstance(c.value, str):
            norm = c.value.strip().lower().replace("_", " ")
            if norm in BANNED or "orcid" in norm or "email" in norm or "e-mail" in norm or "abstract" in norm:
                banned_hdr.append(f"{ws.title}:{c.value}")
            if "author" in norm:
                author_hdr.add(c.value)
add("C06c", "no e-mail or ORCID shapes in any workbook cell",
    "0 hits", f"{ncells} cells in {len(wb.worksheets)} sheets, {len(hits)} hits", not hits, short(hits, 3))
add("C06d", "no workbook column named authors, author full names, abstract, email or orcid",
    "0 such columns", f"{len(banned_hdr)} found", not banned_hdr,
    f"author-related headers seen: {short(sorted(author_hdr), 5)}")
add("C06e", "no workbook cell over 30 words outside README",
    "0 cells", f"{len(long_cells)} cells", not long_cells, short(long_cells, 3))

# ---------------------------------------------------------------- C07 render log
rj = json.loads(guard(OUT / "run_logs" / "render.json").read_text(encoding="utf-8"))
r1 = read_dicts(OUT / "tables" / "R1_render_issues.csv")
od = rj["outputs"]["docx"]
docx = OUT / "render" / "main_v2.docx"
dsize = docx.stat().st_size if docx.is_file() else -1
add("C07a", "outputs.docx exists, bytes > 0, and render/main_v2.docx has that size",
    f"exists {od.get('exists')}, {od.get('bytes')} bytes", f"file present {docx.is_file()}, {dsize} bytes",
    bool(od.get("exists")) and od.get("bytes", 0) > 0 and dsize == od.get("bytes"), f"render.json file field: {od.get('file')}")

kinds = Counter(r["kind"] for r in r1)
ux, ck = rj.get("unresolved_cross_references"), rj.get("citation_keys_missing")
ok = (isinstance(ux, list) and isinstance(ck, list)
      and len(ux) == kinds.get("unresolved cross-reference", 0)
      and len(ck) == kinds.get("citation key not in references.bib", 0))
add("C07b", "render.json reference and citation lists match R1 rows of those kinds",
    f"lists of {len(ux) if isinstance(ux, list) else 'n/a'} and {len(ck) if isinstance(ck, list) else 'n/a'}",
    f"R1 rows: {kinds.get('unresolved cross-reference', 0)} and {kinds.get('citation key not in references.bib', 0)}",
    ok, "both are lists")

ik = rj.get("issue_kinds", {})
add("C07c", "issues_in_R1 equals the number of R1 rows",
    f"issues_in_R1 {rj.get('issues_in_R1')}", f"{len(r1)} R1 rows",
    rj.get("issues_in_R1") == len(r1),
    "issue_kinds per kind also equal R1" if dict(kinds) == ik else f"issue_kinds differ from R1 kinds: {short(set(ik.items()) ^ set(kinds.items()), 2)}")

with zipfile.ZipFile(guard(docx)) as z:
    zn = z.namelist()
has_doc = "word/document.xml" in zn
media = [n for n in zn if n.startswith("word/media/") and not n.endswith("/")]
add("C07d", "docx holds word/document.xml and word/media count equals docx.media_files",
    f"document.xml present; {rj['docx']['media_files']} media files", f"document.xml {has_doc}; {len(media)} media files",
    has_doc and len(media) == rj["docx"]["media_files"],
    "media by extension: " + ", ".join(f"{k} {v}" for k, v in Counter(Path(m).suffix for m in media).items()))

# ---------------------------------------------------------------- write
with open(RESULT, "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["check_id", "what", "expected_from_tables", "recomputed", "verdict", "note"])
    w.writerows(RESULTS)

n_pass = sum(1 for r in RESULTS if r[4] == "PASS")
n_fail = len(RESULTS) - n_pass
print(f"rows {len(RESULTS)}: PASS {n_pass}, FAIL {n_fail}")
print(f"C02 sampled sheets (seed 0): {picked}")
for r in RESULTS:
    if r[4] == "FAIL":
        print("FAIL", r[0], "|", r[1], "|", r[3], "|", r[5])
